from pathlib import Path
import zipfile
import re
from docx.opc.exceptions import PackageNotFoundError
from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError


class ResumeParseError(ValueError):
    """Raised when an uploaded document cannot be safely read as resume text."""


ALLOWED_EXTENSIONS = {".pdf", ".docx"}


def parse_resume(file_path):
    path = Path(file_path)
    extension = path.suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ResumeParseError("Upload a PDF or DOCX resume.")

    try:
        if extension == ".pdf":
            text = _extract_pdf(path)
        else:
            text = _extract_docx(path)
    except ResumeParseError:
        raise
    except (OSError, ValueError, zipfile.BadZipFile, PackageNotFoundError, PdfReadError) as error:
        raise ResumeParseError("This document is unreadable or damaged. Please try another file.") from error

    cleaned = re.sub(r"[ \t\r\f\v]+", " ", text)
    cleaned = re.sub(r"\n[ \t]*\n+", "\n", cleaned).strip()
    if len(cleaned) < 40:
        raise ResumeParseError(
            "No usable text was found. If this is a scanned PDF, this version cannot read it because OCR is not supported."
        )
    return cleaned


def _extract_pdf(path):
    with path.open("rb") as source:
        if source.read(5) != b"%PDF-":
            raise ResumeParseError("The selected file is not a valid PDF.")
        source.seek(0)
        reader = PdfReader(source)
        if reader.is_encrypted:
            raise ResumeParseError("This PDF is password-protected. Upload an unlocked copy.")
        try:
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except (PdfReadError, ValueError, KeyError, TypeError) as error:
            raise ResumeParseError("Text could not be extracted from this PDF.") from error


def _extract_docx(path):
    if not zipfile.is_zipfile(path):
        raise ResumeParseError("The selected file is not a valid DOCX document.")
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        if (
            len(members) > 2000
            or sum(member.file_size for member in members) > 25 * 1024 * 1024
            or "word/document.xml" not in archive.namelist()
        ):
            raise ResumeParseError("This DOCX document is too complex or does not contain a readable document.")
    document = Document(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(paragraphs)
