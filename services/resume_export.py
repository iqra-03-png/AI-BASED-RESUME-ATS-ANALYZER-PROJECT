from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from services.product import RESUME_FIELDS, TEMPLATE_SECTION_ORDER


def _section_text(value):
    return [line.strip().lstrip("•- ").strip() for line in value.splitlines() if line.strip()]


def generate_resume_export(data, template, file_format):
    template_name = template["name"]
    accent_colors = {
        "blue": "176B87", "green": "27735E", "slate": "526675",
        "violet": "7159AD", "cyan": "187C8B", "gold": "9A7137",
        "rose": "A24F6A", "indigo": "4E5AA6",
    }
    accent = accent_colors[template["accent"]]
    section_order = TEMPLATE_SECTION_ORDER[template["id"]]
    if file_format == "docx":
        document = Document()
        title = document.add_heading(data.get("full_name") or "Resume", 0)
        title.alignment = 0 if template["id"] in {"free-classic", "premium-engineering", "premium-hr"} else 1
        title_style = document.styles["Title"]
        title_style.font.color.rgb = RGBColor.from_string(accent)
        title_style.font.size = Pt(24 if template["id"] == "free-compact" else 28)
        heading_style = document.styles["Heading 1"]
        heading_style.font.color.rgb = RGBColor.from_string(accent)
        heading_style.font.size = Pt(11 if template["id"] == "free-compact" else 13)
        if data.get("headline"):
            document.add_paragraph(data["headline"])
        contact = " | ".join(
            data.get(field, "") for field in ("email", "phone", "location") if data.get(field)
        )
        if contact:
            document.add_paragraph(contact)
        for field in section_order:
            label = RESUME_FIELDS[field][0]
            text = data.get(field, "").strip()
            if text:
                document.add_heading(label, level=1)
                for line in _section_text(text):
                    document.add_paragraph(line, style="List Bullet")
        output = BytesIO()
        document.save(output)
        output.seek(0)
        return output

    output = BytesIO()
    pdf = SimpleDocTemplate(
        output, pagesize=A4, rightMargin=20 * mm, leftMargin=20 * mm,
        topMargin=18 * mm, bottomMargin=18 * mm, title=f"{data.get('full_name') or 'Resume'} - {template_name}",
    )
    styles = getSampleStyleSheet()
    title_alignment = TA_CENTER if template["id"] not in {
        "free-classic", "premium-engineering", "premium-hr",
    } else 0
    styles.add(ParagraphStyle(
        name="ResumeName", parent=styles["Title"], alignment=title_alignment,
        textColor=colors.HexColor(f"#{accent}"), fontSize=22 if template["id"] == "free-compact" else 27,
        spaceAfter=5,
    ))
    styles.add(ParagraphStyle(
        name="ResumeSection", parent=styles["Heading2"],
        textColor=colors.HexColor(f"#{accent}"),
        fontSize=10 if template["id"] == "free-compact" else 12,
        spaceBefore=10, spaceAfter=4,
    ))
    content = [Paragraph(escape(data.get("full_name") or "Resume"), styles["ResumeName"])]
    if data.get("headline"):
        content.append(Paragraph(escape(data["headline"]), styles["Heading2"]))
    contact = " | ".join(
        data.get(field, "") for field in ("email", "phone", "location") if data.get(field)
    )
    if contact:
        content.append(Paragraph(escape(contact), styles["BodyText"]))
    content.append(Spacer(1, 4 * mm))
    for field in section_order:
        label = RESUME_FIELDS[field][0]
        text = data.get(field, "").strip()
        if not text:
            continue
        content.append(Paragraph(escape(label), styles["ResumeSection"]))
        for line in _section_text(text):
            content.append(Paragraph(f"• {escape(line)}", styles["BodyText"]))
    pdf.build(content)
    output.seek(0)
    return output
